#!/usr/bin/env python3
"""新建求职工作区：放入配置、建立交接目录、生成空账本与主数据。已有的文件一律不覆盖。

用法：python3 初始化工作区.py --workspace <工作区目录> --skills <campus-job-application 所在的 skills 目录>
例：python3 初始化工作区.py --workspace ~/求职 --skills ~/求职/.agents/skills
"""
import argparse
import json
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
REGION_HEAD = ['地区', '公司', '主要行业', '企业性质', '岗位名称', '工作地点', '岗位JD', '所属BU', '岗位链接', '投递状态',
               '校招计划', '备注', '内推码', '数据刷新时间', '投递截止']
OTHER_SHEETS = {
    '校招主链接': ['公司', '地区', '校招主链接', '适用届别', '投递窗口'],
    '校招计划': ['公司', '计划', '投递关系', '地区', '适用届别', '计划链接', '依据链接', '核验日期', '盘点状态'],
    '投递名额': ['公司', '范围', '适用计划/批次', '投递上限', '当前投递数量', '是否分BU可平行投递', '是否可以滚动投递', '规则原文', '来源/核验日期'],
}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--workspace', required=True)
    ap.add_argument('--skills', required=True, help='含 campus-job-application 的 skills 目录')
    a = ap.parse_args()
    root = Path(a.workspace).expanduser().resolve()
    scripts = Path(a.skills).expanduser().resolve() / 'campus-job-application' / 'scripts'
    if not (scripts / 'json_excel_sync.py').exists():
        sys.exit(f'找不到搜索组件脚本：{scripts}')
    sys.path.insert(0, str(scripts))
    import workbook_store

    done = []
    for src in sorted((HERE / '产品管理').glob('*.json')):
        dst = root / '产品管理' / src.name
        if not dst.exists():
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dst)
            done.append(str(dst.relative_to(root)))
    cfg = json.loads((root / '产品管理/投递状态标签.json').read_text(encoding='utf-8'))
    (root / cfg['companyStore']['dir']).mkdir(parents=True, exist_ok=True)

    book = root / cfg['workbook']
    if not book.exists():
        wb = Workbook()
        wb.remove(wb.active)
        for name in cfg['sheets']:
            wb.create_sheet(name).append(REGION_HEAD)
        for name, head in OTHER_SHEETS.items():
            wb.create_sheet(name).append(head)
        wb.save(book)
        done.append(book.name)

    master = root / cfg['masterData']
    if not master.exists():
        snap = workbook_store.snapshot(book)
        snap['excel_written'] = {'data_sha256': snap['data_sha256'], 'workbook_sha256': snap['workbook_sha256'], 'written_at': None}
        master.parent.mkdir(parents=True, exist_ok=True)
        master.write_text(json.dumps(snap, ensure_ascii=False, indent=2), encoding='utf-8')
        done.append(str(master.relative_to(root)))
    print(json.dumps({'workspace': str(root), 'created': done}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
