"""Synthetic MCP data-flow tests; no real workbook or human consent is used."""
import asyncio
import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import company_handoff as h
import workbook_store as storage
from recruiting_data import RecruitingData
from search_mcp_client import call
from test_company_handoff import sample_company


class RecruitingDataTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.private = self.root / 'private'
        self.private.mkdir()
        self.book = self.root / 'tracker.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = '地区岗位'
        ws.append(['公司', '岗位名称', '所属BU', '工作地点', '岗位链接', '岗位JD', '投递状态'])
        ws.append(['合成公司', 'AI产品', 'BU-A', '城市甲', 'https://example.test/job/1', '职责：AI产品需求。要求：本科。', '待投递'])
        wb.save(self.book)
        self.company = sample_company(self.book)
        self.service = RecruitingData(self.root, self.private)

    def tearDown(self):
        self.tmp.cleanup()

    def put(self, data=None):
        return self.service.draft_put('tracker.xlsx', 'synthetic', data or {'companies': [self.company]})

    def commit(self, draft):
        return self.service.commit('tracker.xlsx', 'synthetic', draft['draft_ref'], draft['draft_sha256'], h.sha(self.book))

    def test_readonly_open_does_not_initialize_or_mutate(self):
        before = h.sha(self.book)
        opened = self.service.open('tracker.xlsx', 'synthetic')
        self.assertEqual(opened['sheets'][0]['company_count'], 1)
        self.assertFalse(opened['full_json_exists'])
        self.assertEqual(h.sha(self.book), before)

    def test_draft_saved_without_excel_mutation_and_tamper_rejected(self):
        before = h.sha(self.book)
        draft = self.put()
        self.assertEqual(h.sha(self.book), before)
        (self.private / draft['draft_ref']).write_text('{}', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '草稿已变化'):
            self.commit(draft)

    def test_foreign_paths_and_batch_drafts_rejected(self):
        with self.assertRaises(ValueError):
            self.service.open(str(self.book), 'synthetic')
        draft = self.put()
        with self.assertRaises(ValueError):
            self.service.draft('tracker.xlsx', 'other', draft['draft_ref'], draft['draft_sha256'])
        with self.assertRaises(ValueError):
            self.service.draft('tracker.xlsx', 'synthetic', '../tracker.xlsx', draft['draft_sha256'])

    def test_stale_write_and_incomplete_confirmation_rejected(self):
        draft = self.put()
        with self.assertRaisesRegex(ValueError, '版本已变化'):
            self.service.commit('tracker.xlsx', 'synthetic', draft['draft_ref'], draft['draft_sha256'], '0' * 64)
        self.commit(draft)
        with self.assertRaises(ValueError):
            self.service.confirm('tracker.xlsx', 'synthetic', 'synthetic', 'missing.md',
                                 h.sha(h.paths(self.book, 'synthetic') / '公司核对.html'), h.sha(self.book))

    def test_full_json_and_company_query_keep_original_rows(self):
        self.commit(self.put())
        self.assertEqual(storage.check_full(self.book), [])
        rows = self.service.query('tracker.xlsx', company='合成公司')['rows']
        self.assertTrue(any(x['sheet'] == '地区岗位' for x in rows))
        self.assertTrue(any(x['sheet'] == '岗位目录' for x in rows))
        with self.assertRaises(ValueError):
            self.service.query('tracker.xlsx')

    def test_existing_submission_is_protected(self):
        wb = load_workbook(self.book)
        wb['地区岗位']['G2'] = '已提交'
        wb.save(self.book)
        draft = self.put({'companies': [self.company], 'workbook_updates': [{
            'sheet': '地区岗位', 'key': {'公司': '合成公司', '岗位链接': 'https://example.test/job/1'},
            'values': {'投递状态': '待投递'}}]})
        with self.assertRaises(ValueError):
            self.commit(draft)
        self.assertEqual(load_workbook(self.book)['地区岗位']['G2'].value, '已提交')

    def test_real_mcp_initialization_draft_excel_review_confirm_adjust(self):
        env = {'CAMPUS_SEARCH_PRIVATE_DIR': str(self.private), 'CAMPUS_RECRUITING_WORKSPACE': str(self.root)}
        def invoke(tool, **arguments):
            return asyncio.run(call({'tool': tool, 'arguments': arguments}))
        common = {'workbook': 'tracker.xlsx', 'batch': 'synthetic'}
        with patch.dict(os.environ, env):
            tools = asyncio.run(call({'method': 'tools/list'}))
            names = {t['name'] for t in tools['tools']}
            self.assertTrue({'search_dispatch', 'recruiting_data_commit', 'recruiting_review_confirm'} <= names)
            opened = invoke('recruiting_data_open', **common)
            invoke('recruiting_data_initialize', **common, sheet='地区岗位', scope='合成测试范围', expected_sha256=opened['workbook_sha256'])
            draft = invoke('recruiting_draft_put', **common, data={'companies': [self.company]})
            invoke('recruiting_data_commit', **common, draft_ref=draft['draft_ref'], draft_sha256=draft['draft_sha256'], expected_sha256=draft['workbook_sha256'])
            html = h.paths(self.book, 'synthetic') / '公司核对.html'
            self.assertIn('AI产品', html.read_text(encoding='utf-8'))
            self.assertEqual(storage.check_full(self.book), [])
            self.assertFalse(invoke('recruiting_handoff_check', **common)['passed'])
            (self.root / 'synthetic-confirm.md').write_text('Synthetic test approval only; no real user authorization.', encoding='utf-8')
            invoke('recruiting_review_confirm', **common, company_key='synthetic', evidence_ref='synthetic-confirm.md', html_sha256=h.sha(html), expected_sha256=h.sha(self.book))
            self.assertTrue(invoke('recruiting_handoff_check', **common, url=self.company['roles'][0]['url'])['passed'])
            changed = copy.deepcopy(self.company)
            changed.update(version='v2', recommendation_reason='合成用户调整后的说明')
            draft = invoke('recruiting_draft_put', **common, data={'companies': [changed]})
            invoke('recruiting_data_commit', **common, draft_ref=draft['draft_ref'], draft_sha256=draft['draft_sha256'], expected_sha256=draft['workbook_sha256'])
            self.assertFalse(invoke('recruiting_handoff_check', **common)['passed'])
            self.assertEqual(storage.check_full(self.book), [])


if __name__ == '__main__':
    unittest.main()
