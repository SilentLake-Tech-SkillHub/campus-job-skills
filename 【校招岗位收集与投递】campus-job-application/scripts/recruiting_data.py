"""Private, version-bound MCP adapters for the existing workbook/JSON/HTML code.

Draft JSON is an exchange input until committed. Excel remains authoritative.
No method in this module accesses a recruiting account or submits an application.
"""
import copy
import json
import os
from contextlib import contextmanager
from pathlib import Path

import company_handoff as handoff
import workbook_store as storage
from openpyxl import load_workbook


class RecruitingData:
    def __init__(self, workspace, private_dir):
        self.workspace = Path(workspace).resolve(strict=True)
        self.private = Path(private_dir).resolve(strict=True)
        if not self.private.is_relative_to(self.workspace):
            raise ValueError('MCP私有目录必须在已配置工作区内')

    def book(self, relative):
        if Path(relative).is_absolute():
            raise ValueError('工作簿使用相对工作区路径')
        p = (self.workspace / relative).resolve(strict=True)
        if not p.is_relative_to(self.workspace) or p.suffix.lower() != '.xlsx':
            raise ValueError('工作簿必须是工作区内的xlsx文件')
        return p

    def batch_path(self, book, batch):
        out = handoff.paths(book, batch)
        if not out.resolve().is_relative_to(self.workspace):
            raise ValueError('批次输出不能通过符号链接越出工作区')
        return out

    @contextmanager
    def lock(self, book, expected_sha256):
        """Serial MCP mutations; existing writer also detects external Excel edits."""
        lock = self.private / ('workbook-' + handoff.digest(str(book))[:24] + '.lock')
        fd = os.open(lock, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        try:
            os.write(fd, str(os.getpid()).encode())
            if handoff.sha(book) != expected_sha256:
                raise ValueError('工作簿版本已变化，重新展开并合并当前数据')
            yield
        finally:
            os.close(fd)
            lock.unlink()

    def open(self, workbook, batch):
        """Small initialization read; never creates sheets or rewrites Excel."""
        book = self.book(workbook)
        self.batch_path(book, batch)
        before = handoff.sha(book)
        wb = load_workbook(book, read_only=True, data_only=False)
        try:
            sheets = []
            for ws in wb:
                headers = list(next(ws.values, ()))
                names = []
                if '公司' in headers:
                    col = headers.index('公司')
                    names = list(dict.fromkeys(r[col] for r in ws.iter_rows(min_row=2, values_only=True) if r[col]))
                sheets.append({'name': ws.title, 'rows': ws.max_row, 'columns': headers,
                               'company_count': len(names)})
        finally:
            wb.close()
        if before != handoff.sha(book):
            raise ValueError('展开期间工作簿变化，请重读')
        full = storage.full_path(book)
        return {'workbook': workbook, 'batch': batch, 'workbook_sha256': before,
                'sheets': sheets, 'full_json_exists': full.exists(),
                'full_json_errors': storage.check_full(book),
                'schema': {kind: dict(fields) for kind, (_, fields) in handoff.SCHEMA.items()},
                'initialization_is_research': False}

    def query(self, workbook, company=None, sheet=None, offset=0, limit=30):
        if not company and not sheet:
            raise ValueError('指定公司或工作表后分页展开')
        if type(offset) is not int or offset < 0 or type(limit) is not int or not 1 <= limit <= 200:
            raise ValueError('分页范围无效')
        book = self.book(workbook)
        rows = storage.query(book, company=company, sheet=sheet)
        return {'workbook_sha256': handoff.sha(book), 'total': len(rows),
                'offset': offset, 'rows': rows[offset:offset + limit]}

    def draft_put(self, workbook, batch, data):
        book = self.book(workbook)
        self.batch_path(book, batch)
        if not isinstance(data, dict) or set(data) - {'companies', 'company_patches', 'workbook_updates'}:
            raise ValueError('草稿须符合固定写入协议')
        if 'companies' in data and 'company_patches' in data:
            raise ValueError('companies与company_patches互斥')
        raw = json.dumps(data, ensure_ascii=False, allow_nan=False).encode('utf-8')
        identity = handoff.digest([str(book), batch])[:24]
        path = self.private / 'drafts' / identity / (handoff.digest(data) + '.json')
        if not path.resolve().is_relative_to(self.private):
            raise ValueError('草稿路径越出私有目录')
        if path.exists() and path.read_bytes() != raw:
            raise ValueError('同摘要草稿内容冲突')
        handoff.atomic(path, raw)
        return {'draft_ref': str(path.relative_to(self.private)), 'draft_sha256': handoff.sha(path),
                'workbook_sha256': handoff.sha(book), 'status': 'exchange_draft',
                'business_validation': 'pending_commit_validation', 'written_to_excel': False}

    def draft(self, workbook, batch, draft_ref, draft_sha256):
        book = self.book(workbook)
        identity = handoff.digest([str(book), batch])[:24]
        path = handoff.safe_ref(self.private, draft_ref)
        if not path.is_relative_to(self.private / 'drafts' / identity):
            raise ValueError('草稿不属于此工作簿和批次')
        if handoff.sha(path) != draft_sha256:
            raise ValueError('草稿已变化')
        return json.loads(path.read_text(encoding='utf-8'))

    def artifacts(self, book, batch):
        out = self.batch_path(book, batch)
        items = [book, storage.full_path(book), out / '公司方案.json', out / '公司核对.html']
        return [{'path': str(p.relative_to(self.workspace)), 'sha256': handoff.sha(p)}
                for p in items if p.is_file()]

    def initialize(self, workbook, batch, sheet, scope, expected_sha256, company=None):
        book = self.book(workbook)
        self.batch_path(book, batch)
        with self.lock(book, expected_sha256):
            result = storage.initialize(book, batch, sheet, scope, company=company)
            errors = storage.check_full(book)
            if errors:
                raise ValueError('; '.join(errors))
            return {**result, 'artifacts': self.artifacts(book, batch), 'user_confirmed': False}

    def commit(self, workbook, batch, draft_ref, draft_sha256, expected_sha256):
        book = self.book(workbook)
        self.batch_path(book, batch)
        data = self.draft(workbook, batch, draft_ref, draft_sha256)
        with self.lock(book, expected_sha256):
            result = handoff.write(book, batch, copy.deepcopy(data))
            errors = storage.check_full(book)
            if errors:
                raise ValueError('; '.join(errors))
            return {'workbook_sha256': result['workbook_sha256'],
                    'companies': len(result['companies']), 'artifacts': self.artifacts(book, batch),
                    'audit': handoff.audit(book, batch), 'application_authorized': False}

    def refresh(self, workbook, batch, expected_sha256):
        book = self.book(workbook)
        self.batch_path(book, batch)
        with self.lock(book, expected_sha256):
            result = storage.refresh(book, batch)
            errors = storage.check_full(book)
            if errors:
                raise ValueError('; '.join(errors))
            return {**result, 'artifacts': self.artifacts(book, batch), 'audit': handoff.audit(book, batch)}

    def confirm(self, workbook, batch, company_key, evidence_ref, html_sha256, expected_sha256):
        """Host may call only after explicit user confirmation of this exact HTML.

        This endpoint cannot manufacture user consent: it requires an existing
        evidence file, and the underlying gate checks research and ranking proof.
        """
        book = self.book(workbook)
        self.batch_path(book, batch)
        with self.lock(book, expected_sha256):
            handoff.confirm(book, batch, company_key, evidence_ref, html_sha256)
            return {'artifacts': self.artifacts(book, batch), 'application_authorized': False}

    def check(self, workbook, batch, url=None):
        book = self.book(workbook)
        self.batch_path(book, batch)
        errors = handoff.check(book, batch, url=url)
        return {'passed': not errors, 'errors': errors, 'application_authorized': False}
