import tempfile,unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook
from workbook_store import apply_updates

class HyperlinkInsertionTests(unittest.TestCase):
 def test_other_company_hyperlink_follows_inserted_row(self):
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);book=root/'sample.xlsx';wb=Workbook();ws=wb.active;ws.title='岗位'
   ws.append(['公司','岗位名称','岗位链接']);ws.append(['甲','原岗','https://example.test/a']);ws.append(['乙','原岗','https://example.test/b']);ws['C3'].hyperlink='https://example.test/b';ws.append([None,None,None])
   apply_updates(wb,[{'sheet':'岗位','key':{'公司':'甲','岗位链接':'https://example.test/a2'},'values':{'岗位名称':'新增'}}],root)
   wb.save(book);wb.close();saved=load_workbook(book);s=saved['岗位']
   self.assertEqual(s['C3'].value,'https://example.test/a2');self.assertIsNone(s['C3'].hyperlink)
   self.assertEqual(s['B4'].value,'原岗');self.assertEqual(s['C4'].hyperlink.target,'https://example.test/b');self.assertEqual(s['C4'].hyperlink.ref,'C4')
   saved.close()
if __name__=='__main__':unittest.main()
