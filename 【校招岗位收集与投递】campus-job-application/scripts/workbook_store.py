#!/usr/bin/env python3
"""从 Excel 生成主数据快照（只用于测试和首次建库）。

日常数据方向是 JSON → Excel：主数据 岗位交接/工作簿全量.json 只经 json_excel_sync.py 写入，
Excel 由它写出。本模块不写 Excel、不写主数据文件。
"""
import hashlib
import json
from datetime import date, datetime, time
from pathlib import Path
from openpyxl import load_workbook

FULL_SCHEMA = 2


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def digest(obj):
    return hashlib.sha256(json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode()).hexdigest()


def value(v):
    if isinstance(v, (datetime, date, time)):
        return {'type': type(v).__name__, 'value': v.isoformat()}
    return v


def snapshot(book):
    """全部工作表与单元格（含公式、缓存值、批注和超链接，不含样式）。"""
    before = sha(book)
    wb = load_workbook(book, data_only=False)
    cached = load_workbook(book, data_only=True)
    sheets = []
    try:
        for ws in wb:
            rows = []
            for row in ws:
                cells = []
                for cell in row:
                    if cell.value is None and not getattr(cell, 'hyperlink', None) and not getattr(cell, 'comment', None):
                        continue
                    item = {'column': cell.column, 'value': value(cell.value), 'type': cell.data_type}
                    if cell.data_type == 'f':
                        item['cached_value'] = value(cached[ws.title][cell.coordinate].value)
                    if cell.hyperlink:
                        item['hyperlink'] = {'target': cell.hyperlink.target, 'location': cell.hyperlink.location}
                    if cell.comment:
                        item['note'] = {'text': cell.comment.text, 'author': cell.comment.author}
                    cells.append(item)
                if cells:
                    rows.append({'row': row[0].row, 'cells': cells})
            sheets.append({'name': ws.title, 'rows': rows, 'max_row': ws.max_row,
                           'max_column': ws.max_column, 'state': ws.sheet_state,
                           'merged_ranges': [str(x) for x in ws.merged_cells.ranges],
                           'tables': [{'name': t.name, 'ref': t.ref} for t in ws.tables.values()]})
    finally:
        wb.close()
        cached.close()
    if before != sha(book):
        raise ValueError('工作簿在快照期间变化')
    return {'schema_version': FULL_SCHEMA, 'scope': 'all_workbook_business_data',
            'workbook_name': Path(book).name, 'workbook_sha256': before,
            'sheets': sheets, 'company_batches': {}, 'data_sha256': digest({'sheets': sheets, 'company_batches': {}})}
