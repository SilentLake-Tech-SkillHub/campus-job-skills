import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

HERE = Path(__file__).resolve().parent
SCRIPT = HERE / 'json_excel_sync.py'
sys.path.insert(0, str(HERE))
import workbook_store  # noqa: E402
import json_excel_sync as S  # noqa: E402

HEAD = ['地区', '公司', '主要行业', '企业性质', '岗位名称', '工作地点', '岗位JD', '所属BU', '岗位链接', '投递状态', '备注']
LABELS = [
    {"value": "岗位待筛选", "definition": "x"},
    {"value": "岗位待投递", "definition": "x"},
    {"value": "岗位已投递", "definition": "x"},
    {"value": "岗位不符合用户要求", "definition": "x"},
]
COMPANY_LABELS = [
    {"value": "待搜索", "definition": "x"},
    {"value": "已有岗位", "definition": "x"},
]


class SyncTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        (self.root / '产品管理').mkdir()
        (self.root / '岗位交接').mkdir()
        wb = Workbook()
        ws = wb.active
        ws.title = '中国大陆'
        ws.append(HEAD)
        ws.append(['上海', '甲公司', '互联网', '私企', None, None, None, None, None, None, None])  # 公司行（方案B：固定保留在公司最前）
        for i in range(4):
            ws.append(['上海', '甲公司', '互联网', '私企', f'产品经理{i}', '上海', 'JD', '', f'https://a.example/{i}', '旧值', ''])
        for r in range(3, 7):
            ws.cell(r, 9).hyperlink = ws.cell(r, 9).value
        ws.append(['北京', '乙公司', '互联网', '私企', None, None, None, None, None, None, None])
        wb.save(self.root / 'book.xlsx')
        (self.root / '产品管理/company-labels.json').write_text(json.dumps(
            {"schemaVersion": 1, "rowKind": "company", "labels": COMPANY_LABELS}, ensure_ascii=False), encoding='utf-8')
        cfg = {"schemaVersion": 1, "rowKind": "role", "labelGroups": ["产品管理/company-labels.json"],
               "column": "投递状态", "sheets": ["中国大陆"], "labels": LABELS,
               "confirmationLedger": "产品管理/确认.json", "masterData": "岗位交接/工作簿全量.json", "workbook": "book.xlsx",
               "quotaRule": {"countsStatuses": ["岗位待投递", "岗位已投递"],
                             "autoCloseFrom": ["岗位待筛选"], "autoCloseTo": "岗位不符合用户要求"}}
        (self.root / '产品管理/labels.json').write_text(json.dumps(cfg, ensure_ascii=False), encoding='utf-8')
        snap = workbook_store.snapshot(self.root / 'book.xlsx')
        (self.root / '岗位交接/工作簿全量.json').write_text(json.dumps(snap, ensure_ascii=False), encoding='utf-8')
        self.run_cmd('set-quota-rules', '--input', self.items([self.quota(2)]))

    def quota(self, limit, company='甲公司', scope='公司总量', text='每人限投2个岗位', source=None):
        src = self.root / ('src-%d.txt' % len(list(self.root.glob('src-*.txt'))))
        src.write_text(source if source is not None else '校招FAQ\n问：可以投几个？\n答：' + text + '。\n', encoding='utf-8')
        return {"公司": company, "范围": scope, "适用计划/批次": "全部计划", "投递上限": limit,
                "是否分BU可平行投递": "公司未披露", "是否可以滚动投递": "否", "规则原文": text,
                "来源/核验日期": f"来源：{company}-校招FAQ 核验日期：2026-10-20", "原文出处文件": str(src)}

    def tearDown(self):
        self.tmp.cleanup()

    def run_cmd(self, *args, ok=True):
        p = subprocess.run([sys.executable, '-B', str(SCRIPT), '--workspace', str(self.root), '--labels', '产品管理/labels.json', *args],
                           capture_output=True, text=True, encoding='utf-8')
        if ok and p.returncode:
            self.fail(p.stderr + p.stdout)
        return p

    def items(self, rows):
        f = self.root / ('in-%d.json' % len(list(self.root.glob('in-*.json'))))
        f.write_text(json.dumps(rows, ensure_ascii=False), encoding='utf-8')
        return str(f)

    def role(self, i, to):
        return {"sheet": "中国大陆", "company": "甲公司", "link": f"https://a.example/{i}", "title": f"产品经理{i}", "to": to}

    def migrate(self):
        rows = [self.role(i, '岗位待筛选') for i in range(4)] + [{"sheet": "中国大陆", "company": "乙公司", "to": "待搜索"},
                                                                {"sheet": "中国大陆", "company": "甲公司", "to": "已有岗位"}]
        self.run_cmd('confirm', '--input', self.items(rows), '--quote', '确认迁移')

    def test_illegal_or_empty_label_blocks_write(self):
        p = self.run_cmd('write', ok=False)
        self.assertNotEqual(p.returncode, 0)
        self.assertIn('非法标签', p.stderr)
        self.assertIn('投递状态为空', p.stderr)

    def test_unknown_label_and_wrong_row_kind_rejected(self):
        p = self.run_cmd('propose', '--input', self.items([self.role(0, '已投递了')]), ok=False)
        self.assertIn('非法标签', p.stderr)
        p = self.run_cmd('propose', '--input', self.items([{"sheet": "中国大陆", "company": "乙公司", "to": "岗位待筛选"}]), ok=False)
        self.assertIn('不适用', p.stderr)

    def test_unconfirmed_status_edit_in_json_blocks_write(self):
        self.migrate()
        self.run_cmd('write')
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        sheet = master['sheets'][0]
        cell = next(c for c in sheet['rows'][2]['cells'] if c['column'] == 10)  # 第3行=第一个岗位行
        cell['value'] = '岗位待投递'  # 绕过脚本直接改 JSON
        (self.root / '岗位交接/工作簿全量.json').write_text(json.dumps(master, ensure_ascii=False), encoding='utf-8')
        p = self.run_cmd('write', ok=False)
        self.assertIn('当前状态是「岗位待筛选」，更新状态是「岗位待投递」', p.stderr)

    def test_propose_shows_current_and_new_and_changes_nothing(self):
        before = (self.root / '岗位交接/工作簿全量.json').read_bytes()
        out = json.loads(self.run_cmd('propose', '--input', self.items([self.role(0, '岗位待筛选')])).stdout)
        self.assertEqual(out['show_user_exactly'], ['中国大陆｜甲公司｜产品经理0：当前状态是「旧值」，更新状态是「岗位待筛选」'])
        self.assertEqual(before, (self.root / '岗位交接/工作簿全量.json').read_bytes())

    def test_confirm_write_readback_and_dropdown(self):
        self.migrate()
        pend = json.loads(self.run_cmd('pending').stdout)
        self.assertTrue(pend['unwritten'])
        out = json.loads(self.run_cmd('write').stdout)
        self.assertEqual(out['status_changes_applied'], 6)
        ws = load_workbook(self.root / 'book.xlsx')['中国大陆']
        self.assertEqual([ws.cell(r, 10).value for r in range(2, 8)], ['已有岗位'] + ['岗位待筛选'] * 4 + ['待搜索'])
        self.assertTrue(any('岗位待筛选' in dv.formula1 for dv in ws.data_validations.dataValidation))
        self.assertEqual(ws.cell(3, 9).hyperlink.target, 'https://a.example/0')
        self.assertFalse(json.loads(self.run_cmd('pending').stdout)['unwritten'])
        self.assertEqual(json.loads(self.run_cmd('check').stdout)['label_problems'], 0)

    def test_quota_closure_after_user_fills_quota(self):
        self.migrate()
        out = json.loads(self.run_cmd('propose', '--input', self.items([self.role(0, '岗位待投递'), self.role(1, '岗位待投递')])).stdout)
        self.assertEqual(out['quota_closure_changes'], 2)
        self.assertTrue(all('名额收口' in t for t in out['show_user_exactly'][2:]))
        res = json.loads(self.run_cmd('confirm', '--input', self.items([self.role(0, '岗位待投递'), self.role(1, '岗位待投递')]), '--quote', '选0和1').stdout)
        self.assertEqual(len(res['quota_closure_changes']), 2)
        self.run_cmd('write')
        ws = load_workbook(self.root / 'book.xlsx')['中国大陆']
        self.assertEqual([ws.cell(r, 10).value for r in range(3, 7)], ['岗位待投递', '岗位待投递', '岗位不符合用户要求', '岗位不符合用户要求'])

    def test_no_closure_before_quota_is_filled_and_overfill_rejected(self):
        self.migrate()
        out = json.loads(self.run_cmd('propose', '--input', self.items([self.role(0, '岗位待投递')])).stdout)
        self.assertEqual(out['quota_closure_changes'], 0)
        p = self.run_cmd('confirm', '--input', self.items([self.role(i, '岗位待投递') for i in range(3)]), '--quote', 'x', ok=False)
        self.assertIn('超过本轮名额', p.stderr)

    def test_upsert_cannot_change_status_and_new_role_needs_confirmation(self):
        self.migrate()
        self.run_cmd('write')
        p = self.run_cmd('upsert', '--input', self.items([{"key": {"sheet": "中国大陆", "company": "乙公司", "link": "https://b.example/1", "title": "AI产品"}, "values": {"投递状态": "岗位待投递"}}]), ok=False)
        self.assertIn('不能修改投递状态', p.stderr)
        self.run_cmd('upsert', '--input', self.items([{"key": {"sheet": "中国大陆", "company": "乙公司", "link": "https://b.example/1", "title": "AI产品"}, "values": {"工作地点": "北京"}}]))
        p = self.run_cmd('write', ok=False)
        self.assertIn('投递状态为空', p.stderr)

    def test_upsert_clearing_link_removes_hyperlink_and_duplicate_company_row_blocked(self):
        self.migrate()
        self.run_cmd('write')
        self.run_cmd('upsert', '--input', self.items([{"key": {"sheet": "中国大陆", "company": "甲公司", "link": "https://a.example/3", "title": "产品经理3"}, "values": {"岗位链接": None, "岗位名称": None}}]))
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        row = next(r for r in master['sheets'][0]['rows'] if r['row'] == 6)
        self.assertFalse(any(c['column'] == 9 for c in row['cells']))
        p = self.run_cmd('write', ok=False)  # 名称和链接都清空后变成第二个公司行：业务键重复，且岗位标签不适用于公司行
        self.assertIn('业务键重复', p.stderr)
        self.assertIn('不适用于公司行', p.stderr)

    def test_clear_orphans_only_for_rows_without_company(self):
        p = self.run_cmd('clear-orphans', '--input', self.items([{"sheet": "中国大陆", "row": 2}]), '--quote', 'x', ok=False)
        self.assertIn('有公司名', p.stderr)
        wb = load_workbook(self.root / 'book.xlsx'); ws = wb['中国大陆']; ws.cell(8, 9).value = 'https://ghost'; wb.save(self.root / 'book.xlsx')
        snap = workbook_store.snapshot(self.root / 'book.xlsx')
        (self.root / '岗位交接/工作簿全量.json').write_text(json.dumps(snap, ensure_ascii=False), encoding='utf-8')
        self.migrate()
        self.run_cmd('clear-orphans', '--input', self.items([{"sheet": "中国大陆", "row": 8}]), '--quote', '清空')
        self.run_cmd('write')
        self.assertIsNone(load_workbook(self.root / 'book.xlsx')['中国大陆'].cell(8, 9).value)

    def test_no_accumulating_snapshots(self):
        self.migrate()
        self.run_cmd('write')
        for city in ('北京', '上海'):
            self.run_cmd('upsert', '--input', self.items([{"key": {"sheet": "中国大陆", "company": "乙公司", "link": "https://b.example/1", "title": "AI产品"}, "values": {"工作地点": city}}]))
        self.assertFalse((self.root / '岗位交接/备份').exists())
        self.assertEqual(sorted(p.name for p in (self.root / '岗位交接').iterdir()), ['工作簿全量.json'])

    def test_quota_rules_single_latest_version_per_company(self):
        """名额表：同一公司再次写入时整体覆盖，只保留最新一版；值只能是 是/否/公司未披露；当前投递数量由脚本计算。"""
        self.migrate()
        self.run_cmd('set-quota-rules', '--input', self.items([self.quota(1, text='新规则：限投1个'),
                                                                self.quota('公司未披露', scope='P-STAR计划', text='独立计划，额度未披露')]))
        self.run_cmd('set-quota-rules', '--input', self.items([self.quota(3, company='乙公司', text='乙限投3个')]))
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        rows = [r for r, v in S.quota_records(master)]
        recs = [v for r, v in S.quota_records(master)]
        self.assertEqual([(v['公司'], v['范围'], v['投递上限']) for v in recs],
                         [('乙公司', '公司总量', 3), ('甲公司', 'P-STAR计划', '公司未披露'), ('甲公司', '公司总量', 1)])
        self.assertNotIn('每人限投2个岗位', json.dumps(master, ensure_ascii=False))
        self.run_cmd('confirm', '--input', self.items([self.role(0, '岗位待投递')]), '--quote', '选0')
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        self.assertEqual({(v['公司'], v['范围']): v['当前投递数量'] for r, v in S.quota_records(master)}[('甲公司', '公司总量')], 1)
        for bad, msg in ((dict(self.quota(1), **{"是否可以滚动投递": "未知"}), '是/否/公司未披露'),
                         (dict(self.quota('未知')), '投递上限'),
                         (dict(self.quota(1), **{"当前投递数量": 5}), '当前投递数量')):
            p = self.run_cmd('set-quota-rules', '--input', self.items([bad]), ok=False)
            self.assertIn(msg, p.stderr)
        p = self.run_cmd('set-quota-rules', '--input', self.items([self.quota(1), self.quota(2)]), ok=False)
        self.assertIn('只能有一条规则', p.stderr)
        self.run_cmd('write')
        ws = load_workbook(self.root / 'book.xlsx')['投递名额']
        self.assertEqual(ws.cell(1, 5).value, '当前投递数量')

    def test_quota_quote_must_be_verbatim_and_source_format_fixed(self):
        """规则原文只能是来源原文的逐字摘抄；来源/核验日期格式由代码卡死。"""
        page = '20XX届校园招聘-正式批每位同学共计3岗。区域公司项目不限。提交后不可修改。'
        ok = self.quota(3, text='20XX届校园招聘-正式批每位同学共计3岗。', source=page)
        self.run_cmd('set-quota-rules', '--input', self.items([ok]))
        bad = [
            (dict(ok, **{"规则原文": "正式批最多投3个岗位"}), '逐字连续摘抄'),                       # 总结改写
            (dict(ok, **{"规则原文": "正式批每位同学共计3岗。；提交后不可修改。"}), '逐字连续摘抄'),    # 拼接多段
            (dict(ok, **{"规则原文": "本维度已查事实见本轮规则汇总，完整关系未齐。"}), '逐字连续摘抄'),  # 过程话
            ({k: v for k, v in ok.items() if k != '原文出处文件'}, '原文出处文件'),
            (dict(ok, **{"来源/核验日期": "官网 2026-10-20"}), '来源：<公司>-<岗位或页面> 核验日期：YYYY-MM-DD'),
            (dict(ok, **{"来源/核验日期": "来源：乙公司-校招FAQ 核验日期：2026-10-20"}), '公司与本行一致'),
            (dict(ok, **{"来源/核验日期": "来源：甲公司-校招FAQ 核验日期：2026-13-01"}), '有效日期'),
            (dict(ok, **{"来源/核验日期": "来源：甲公司-校招FAQ 核验日期：2999-01-01"}), '晚于今天'),
        ]
        for item, msg in bad:
            p = self.run_cmd('set-quota-rules', '--input', self.items([item]), ok=False)
            self.assertIn(msg, p.stderr, item)
        # 没有官方说法：规则原文写“公司未披露”，不需要原文文件，来源格式照样卡死
        none = {k: v for k, v in self.quota('公司未披露', text='公司未披露').items() if k != '原文出处文件'}
        self.run_cmd('set-quota-rules', '--input', self.items([none]))
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        self.assertNotIn('原文出处文件', json.dumps(master, ensure_ascii=False))
        chk = json.loads(self.run_cmd('check', ok=False).stdout)
        self.assertEqual(chk['quota_rows_not_in_current_format'], 0)

    def test_drop_rule_column_keeps_excel_table_valid(self):
        wb = load_workbook(self.root / 'book.xlsx')
        ws = wb.create_sheet('校招计划')
        ws.append(['公司', '计划', '计划内投递限制', '核验日期'])
        ws.append(['甲公司', '常规', '限投1个', '2026-09-25'])
        from openpyxl.worksheet.table import Table
        ws.add_table(Table(displayName='Plans', ref='A1:D2'))
        wb.save(self.root / 'book.xlsx')
        snap = workbook_store.snapshot(self.root / 'book.xlsx')
        (self.root / '岗位交接/工作簿全量.json').write_text(json.dumps(snap, ensure_ascii=False), encoding='utf-8')
        self.run_cmd('set-quota-rules', '--input', self.items([self.quota(2)]))
        self.migrate()
        self.run_cmd('drop-column', '--sheet', '校招计划', '--name', '计划内投递限制')
        self.run_cmd('write')
        ws = load_workbook(self.root / 'book.xlsx')['校招计划']
        self.assertEqual([c.value for c in ws[1]][:3], ['公司', '计划', '核验日期'])
        self.assertEqual(ws.tables['Plans'].ref, 'A1:C2')
        self.assertEqual([c.name for c in ws.tables['Plans'].tableColumns], ['公司', '计划', '核验日期'])

    def test_set_sheet_rows_overwrites_notes_sheet_only(self):
        wb = load_workbook(self.root / 'book.xlsx')
        ws = wb.create_sheet('分类说明')
        for r in [['项目', '说明'], ['旧1', '旧说明1'], ['旧2', '旧说明2'], ['旧3', '旧说明3']]:
            ws.append(r)
        from openpyxl.worksheet.table import Table
        ws.add_table(Table(displayName='Notes', ref='A1:B4'))
        wb.save(self.root / 'book.xlsx')
        snap = workbook_store.snapshot(self.root / 'book.xlsx')
        (self.root / '岗位交接/工作簿全量.json').write_text(json.dumps(snap, ensure_ascii=False), encoding='utf-8')
        self.run_cmd('set-quota-rules', '--input', self.items([self.quota(2)]))
        self.migrate()
        self.run_cmd('set-sheet-rows', '--sheet', '分类说明', '--input', self.items([['项目', '说明'], ['新', '新说明']]))
        for name in ('中国大陆', '投递名额'):
            p = self.run_cmd('set-sheet-rows', '--sheet', name, '--input', self.items([['a']]), ok=False)
            self.assertNotEqual(p.returncode, 0)
        self.run_cmd('write')
        ws = load_workbook(self.root / 'book.xlsx')['分类说明']
        self.assertEqual([[c.value for c in r] for r in ws.iter_rows(min_row=1, max_row=4)],
                         [['项目', '说明'], ['新', '新说明'], [None, None], [None, None]])
        self.assertEqual(ws.tables['Notes'].ref, 'A1:B2')

    def test_query_reads_master_json_only(self):
        out = json.loads(self.run_cmd('query', '--company', '乙公司').stdout)
        self.assertEqual(out['total'], 1)
        self.assertEqual(out['rows'][0]['values']['公司'], '乙公司')
        out = json.loads(self.run_cmd('query', '--sheet', '中国大陆', '--offset', '1', '--limit', '2').stdout)
        self.assertEqual((out['total'], len(out['rows']), out['rows'][0]['values']['岗位名称']), (6, 2, '产品经理0'))
        self.assertNotEqual(self.run_cmd('query', ok=False).returncode, 0)

    def test_diff_pairs_rows_by_identity_not_position(self):
        self.migrate() if hasattr(self, 'migrate') else None
        self.run_cmd('write')
        m = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        rows = m['sheets'][0]['rows']
        # 在第2行前插入一个新岗位：后面所有行下移一行，但内容没变
        for r in rows:
            if r['row'] >= 2:
                r['row'] += 1
        rows.insert(1, {'row': 2, 'cells': [{'column': 1, 'value': '上海', 'type': 's'}, {'column': 2, 'value': '甲公司', 'type': 's'},
                                             {'column': 5, 'value': '新岗位', 'type': 's'}, {'column': 9, 'value': 'https://a.example/new', 'type': 's'}]})
        for r in rows:
            if any(c['column'] == 2 and c['value'] == '乙公司' for c in r['cells']):
                for c in r['cells']:
                    if c['column'] == 3:
                        c['value'] = '金融'
        (self.root / '岗位交接/工作簿全量.json').write_text(json.dumps(m, ensure_ascii=False), encoding='utf-8')
        d = json.loads(self.run_cmd('diff', '--limit', '0').stdout)
        self.assertEqual((d['added'], d['removed'], d['changed'], d['cells']), (1, 0, 1, 1))
        self.assertGreaterEqual(d['moved'], 4)
        self.assertIn('新增：中国大陆｜甲公司｜新岗位', d['text'])
        self.assertIn('新岗位｜投递状态：', d['text'])  # 新增行同时显示投递状态
        self.assertIn('修改：中国大陆｜乙公司｜公司行（第8行）：主要行业：当前「互联网」→写入「金融」', d['text'])
        self.assertNotIn('公司：当前', d['text'])  # 不再把位置挪动报成换公司

    def test_report_summarizes_company_from_master_json(self):
        out = json.loads(self.run_cmd('report', '--company', '甲公司', '--company', '乙公司').stdout)
        a, b = out
        self.assertEqual((a['岗位']['总数'], a['岗位']['按投递状态'], a['岗位']['用户已选（计入名额）']), (4, {'旧值': 4}, 0))
        self.assertEqual([(q['范围'], q['投递上限'], q['剩余']) for q in a['投递名额']], [('公司总量', 2, 2)])
        self.assertEqual((b['岗位']['总数'], b['投递名额'], b['地区表']), (0, [], ['中国大陆']))
        self.assertNotEqual(self.run_cmd('report', '--company', '不存在', ok=False).returncode, 0)

    def test_field_rules_for_remark_and_deadline(self):
        rules = {"schemaVersion": 1, "fields": {
            "备注": {"格式": "^用户例外保留：.{4,}$", "最大长度": 150, "适用行": ["role"], "允许为空": True},
            "投递截止": {"格式": "^(\\d{4}-\\d{2}-\\d{2}|公司未披露)$", "适用行": ["role"], "岗位行必填": True}}}
        (self.root / '产品管理/字段.json').write_text(json.dumps(rules, ensure_ascii=False), encoding='utf-8')
        cfg = json.loads((self.root / '产品管理/labels.json').read_text(encoding='utf-8'))
        cfg['fieldRules'] = '产品管理/字段.json'
        (self.root / '产品管理/labels.json').write_text(json.dumps(cfg, ensure_ascii=False), encoding='utf-8')
        self.migrate()
        p = self.run_cmd('write', ok=False)
        self.assertIn('缺少列「投递截止」', p.stderr)
        self.run_cmd('add-column', '--sheet', '中国大陆', '--name', '投递截止')
        key = lambda i: {"sheet": "中国大陆", "company": "甲公司", "link": f"https://a.example/{i}", "title": f"产品经理{i}"}
        p = self.run_cmd('upsert', '--input', self.items([{"key": key(0), "values": {"备注": "Codex首轮已查，见 岗位交接/x.json"}}]), ok=False)
        self.assertIn('备注', p.stderr)
        p = self.run_cmd('upsert', '--input', self.items([{"key": key(0), "values": {"投递截止": "下周五"}}]), ok=False)
        self.assertIn('投递截止', p.stderr)
        self.run_cmd('upsert', '--input', self.items([{"key": key(i), "values": {"投递截止": "公司未披露"}} for i in range(4)]))
        self.run_cmd('upsert', '--input', self.items([{"key": key(1), "values": {"备注": "用户例外保留：城市在苏州但用户要求保留"}}]))
        self.run_cmd('write')
        ws = load_workbook(self.root / 'book.xlsx')['中国大陆']
        self.assertEqual(ws.cell(1, 12).value, '投递截止')
        self.assertEqual(ws.cell(4, 11).value, '用户例外保留：城市在苏州但用户要求保留')

    def test_title_with_choice_order_rejected(self):
        """一行拼多个岗位、岗位名称里写“第N志愿”：upsert（含业务键带入的名称）、check、write 都必须拦。"""
        rules = {"schemaVersion": 1, "fields": {"岗位名称": {"适用行": ["role"], "允许为空": True,
                 "禁止模式": [{"模式": "志愿", "原因": "含选岗志愿顺序"}]}}}
        (self.root / '产品管理/字段.json').write_text(json.dumps(rules, ensure_ascii=False), encoding='utf-8')
        cfg = json.loads((self.root / '产品管理/labels.json').read_text(encoding='utf-8'))
        cfg['fieldRules'] = '产品管理/字段.json'
        (self.root / '产品管理/labels.json').write_text(json.dumps(cfg, ensure_ascii=False), encoding='utf-8')
        self.migrate()
        bad = 'A岗（第1志愿） / B岗（第2志愿）'
        # 新增行：名称只在业务键里
        p = self.run_cmd('upsert', '--input', self.items([{"key": {"sheet": "中国大陆", "company": "甲公司",
                         "link": "https://a.example/9", "title": bad}, "values": {"工作地点": "上海"}}]), ok=False)
        self.assertIn('志愿', p.stderr)
        # 改名：名称在 values 里
        key0 = {"sheet": "中国大陆", "company": "甲公司", "link": "https://a.example/0", "title": "产品经理0"}
        p = self.run_cmd('upsert', '--input', self.items([{"key": key0, "values": {"岗位名称": bad}}]), ok=False)
        self.assertIn('志愿', p.stderr)
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        self.assertNotIn(bad, [c.get('value') for r in master['sheets'][0]['rows'] for c in r['cells']])
        # 已在主数据里的历史脏行：check 报出、write 拒绝
        S.set_cell(master['sheets'][0]['rows'][1], 5, bad)
        (self.root / '岗位交接/工作簿全量.json').write_text(json.dumps(master, ensure_ascii=False), encoding='utf-8')
        out = json.loads(self.run_cmd('check', ok=False).stdout)
        self.assertEqual(out['field_problems'], 1)
        self.assertIn('志愿', out['field_problem_samples'][0])
        p = self.run_cmd('write', ok=False)
        self.assertIn('志愿', p.stderr)

    def test_rename_is_not_a_status_change(self):
        """改岗位名称（业务键的一部分）后，按同公司同链接找回 Excel 原行，不误报为未确认的状态变更。"""
        self.migrate()
        self.run_cmd('write')
        key0 = {"sheet": "中国大陆", "company": "甲公司", "link": "https://a.example/0", "title": "产品经理0"}
        self.run_cmd('upsert', '--input', self.items([{"key": key0, "values": {"岗位名称": "AI产品经理（官网原名）"}}]))
        out = json.loads(self.run_cmd('check', ok=False).stdout)
        self.assertEqual((out['status_changes'], out['status_changes_without_user_confirmation']), (0, 0))
        self.run_cmd('write')

    def test_reset_rows_back_to_search(self):
        """删除无效岗位行：公司行固定保留；岗位删光时公司行改回“待搜索”并删名额规则；已投递的行不能删；其它行只上移不乱序。"""
        self.migrate()
        self.run_cmd('confirm', '--input', self.items([self.role(3, '岗位已投递')]), '--quote', '已投')
        p = self.run_cmd('reset-rows', '--input', self.items([{k: v for k, v in self.role(3, '').items() if k != 'to'}]), '--quote', 'x', ok=False)
        self.assertIn('真实投递记录', p.stderr)
        rows = [{k: v for k, v in self.role(i, '').items() if k != 'to'} for i in range(3)]
        prev = json.loads(self.run_cmd('reset-rows', '--input', self.items(rows), '--preview').stdout)
        self.assertEqual(len(prev['show_user_exactly']), 3)
        self.run_cmd('reset-rows', '--input', self.items(rows), '--quote', '退回待搜索')
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        sheet = master['sheets'][0]
        names = [(r['row'], S.cell_value(r, 2), S.cell_value(r, 5)) for r in sheet['rows'][1:]]
        self.assertEqual(names, [(2, '甲公司', None), (3, '甲公司', '产品经理3'), (4, '乙公司', None)])
        self.assertEqual([v['公司'] for r, v in S.quota_records(master)], ['甲公司'])
        rows = [{"sheet": "中国大陆", "company": "甲公司", "link": "https://a.example/3", "title": "产品经理3"}]
        self.run_cmd('write')

    def test_reset_whole_company_adds_placeholder(self):
        self.migrate()
        rows = [{k: v for k, v in self.role(i, '').items() if k != 'to'} for i in range(4)]
        self.run_cmd('reset-rows', '--input', self.items(rows), '--quote', '退回待搜索')
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        rows = [(r['row'], S.cell_value(r, 2), S.cell_value(r, 5), S.cell_value(r, 10)) for r in master['sheets'][0]['rows'][1:]]
        self.assertEqual(rows, [(2, '甲公司', None, '待搜索'), (3, '乙公司', None, '待搜索')])
        self.assertEqual(S.quota_records(master), [])
        self.run_cmd('write')

    def test_drop_company_without_placeholder(self):
        self.migrate()
        rows = [{k: v for k, v in self.role(i, '').items() if k != 'to'} for i in range(4)]
        self.run_cmd('reset-rows', '--input', self.items(rows), '--quote', '直接删除', '--drop-company', '甲公司')
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        rows = [(r['row'], S.cell_value(r, 2)) for r in master['sheets'][0]['rows'][1:]]
        self.assertEqual(rows, [(2, '乙公司')])
        self.assertEqual(S.quota_records(master), [])
        self.run_cmd('write')

    def test_mark_research_keeps_other_rows_and_rules(self):
        self.migrate()
        rows = [{k: v for k, v in self.role(0, '').items() if k != 'to'}]
        self.run_cmd('reset-rows', '--input', self.items(rows), '--quote', '重新搜索', '--mark-research')
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        rows = [(r['row'], S.cell_value(r, 2), S.cell_value(r, 5), S.cell_value(r, 10)) for r in master['sheets'][0]['rows'][1:]]
        self.assertEqual(rows[0], (2, '甲公司', None, '待搜索'))
        self.assertEqual([r[2] for r in rows[1:4]], ['产品经理1', '产品经理2', '产品经理3'])
        self.assertEqual([v['公司'] for r, v in S.quota_records(master)], ['甲公司'])
        self.run_cmd('write')

    def relink_item(self, i, new, old=None):
        return {"sheet": "中国大陆", "company": "甲公司", "link": old or f"https://a.example/{i}", "title": f"产品经理{i}", "new_link": new}

    def test_relink_keeps_status_without_reconfirmation(self):
        """换链接：upsert 不能直接改已有行链接；relink 预览不改文件、无原话不执行；确认后写 Excel 状态不变、不要求重新确认。"""
        self.migrate()
        self.run_cmd('confirm', '--input', self.items([self.role(0, '岗位待投递')]), '--quote', '选0')
        self.run_cmd('write')
        key0 = {k: v for k, v in self.role(0, '').items() if k != 'to'}
        p = self.run_cmd('upsert', '--input', self.items([{"key": key0, "values": {"岗位链接": "https://a.example/0-new"}}]), ok=False)
        self.assertIn('请用 relink', p.stderr)
        before = (self.root / '岗位交接/工作簿全量.json').read_bytes()
        out = json.loads(self.run_cmd('relink', '--input', self.items([self.relink_item(0, 'https://a.example/0-new')]), '--preview').stdout)
        self.assertEqual(out['show_user_exactly'], ['中国大陆｜甲公司｜产品经理0：岗位链接「https://a.example/0」→「https://a.example/0-new」'
                                                    '（投递状态「岗位待投递」保持不变，志愿顺序等投递数据随行迁移）'])
        self.assertEqual(before, (self.root / '岗位交接/工作簿全量.json').read_bytes())
        self.assertIn('原话', self.run_cmd('relink', '--input', self.items([self.relink_item(0, 'https://a.example/0-new')]), ok=False).stderr)
        self.run_cmd('relink', '--input', self.items([self.relink_item(0, 'https://a.example/0-new')]), '--quote', '换成官网新链接')
        chk = json.loads(self.run_cmd('check', ok=False).stdout)
        self.assertEqual((chk['link_changes_waiting_for_excel'], chk['status_changes_without_user_confirmation']), (1, 0))
        self.assertEqual(json.loads(self.run_cmd('write').stdout)['status_changes_applied'], 0)
        ws = load_workbook(self.root / 'book.xlsx')['中国大陆']
        self.assertEqual((ws.cell(3, 9).value, ws.cell(3, 9).hyperlink.target, ws.cell(3, 10).value),
                         ('https://a.example/0-new', 'https://a.example/0-new', '岗位待投递'))
        self.assertEqual(json.loads(self.run_cmd('check').stdout)['link_changes_waiting_for_excel'], 0)

    def test_relink_rejects_bad_input_carries_confirmations_and_chains(self):
        """新链接与同公司其它行重复、非 https、公司行都拒绝；未写入的确认随行迁移；写 Excel 前连续换两次只保留一条身份变更。"""
        self.migrate()
        self.run_cmd('write')
        for item, msg in ((self.relink_item(0, 'https://a.example/1'), '同公司其它行使用'),
                          (self.relink_item(0, 'http://a.example/x'), 'https'),
                          ({"sheet": "中国大陆", "company": "乙公司", "new_link": "https://b.example/1"}, '公司行')):
            self.assertIn(msg, self.run_cmd('relink', '--input', self.items([item]), '--preview', ok=False).stderr)
        self.run_cmd('relink', '--input', self.items([self.relink_item(0, 'https://a.example/0-v2')]), '--quote', '第一次')
        self.run_cmd('relink', '--input', self.items([self.relink_item(0, 'https://a.example/0-v3', 'https://a.example/0-v2')]), '--quote', '第二次')
        self.run_cmd('confirm', '--input', self.items([self.role(1, '岗位待投递')]), '--quote', '选1')
        self.run_cmd('relink', '--input', self.items([self.relink_item(1, 'https://a.example/1-new')]), '--quote', '1也换')
        ledger = json.loads((self.root / '产品管理/确认.json').read_text(encoding='utf-8'))
        self.assertEqual(sorted((c['old_key']['link'], c['new_key']['link']) for c in ledger['identity_changes'] if not c.get('applied_at')),
                         [('https://a.example/0', 'https://a.example/0-v3'), ('https://a.example/1', 'https://a.example/1-new')])
        self.assertEqual(json.loads(self.run_cmd('write').stdout)['status_changes_applied'], 1)  # 只有 1 的「待筛选→待投递」
        ws = load_workbook(self.root / 'book.xlsx')['中国大陆']
        self.assertEqual([(ws.cell(r, 9).value, ws.cell(r, 10).value) for r in (3, 4)],
                         [('https://a.example/0-v3', '岗位待筛选'), ('https://a.example/1-new', '岗位待投递')])

    def test_link_exception_note_only_by_command(self):
        """链接例外备注只能由 link-exception 写：upsert 写或改都拒绝；预览不改文件；确认后记入台账；可撤销；绕过脚本写入的被门禁拦下。"""
        self.migrate()
        self.run_cmd('write')
        key0 = {k: v for k, v in self.role(0, '').items() if k != 'to'}
        p = self.run_cmd('upsert', '--input', self.items([{"key": key0, "values": {"备注": S.LINK_EXCEPTION_PREFIX + '列表页按ID定位'}}]), ok=False)
        self.assertIn('link-exception', p.stderr)
        item = {**key0, "定位说明": "列表页按 positionId 定位，点岗位名展开 JD"}
        before = (self.root / '岗位交接/工作簿全量.json').read_bytes()
        out = json.loads(self.run_cmd('link-exception', '--input', self.items([item]), '--preview').stdout)
        self.assertEqual(out['show_user_exactly'], ['中国大陆｜甲公司｜产品经理0：备注「空」→「用户例外保留：官网无独立详情页。列表页按 positionId 定位，点岗位名展开 JD」'
                                                    '（生成投递批次时不检查该岗位链接是否为具体岗位页）'])
        self.assertEqual(before, (self.root / '岗位交接/工作簿全量.json').read_bytes())
        self.run_cmd('link-exception', '--input', self.items([item]), '--quote', '这种可以')
        p = self.run_cmd('upsert', '--input', self.items([{"key": key0, "values": {"备注": "用户例外保留：别的原因写这里"}}]), ok=False)
        self.assertIn('link-exception', p.stderr)
        self.run_cmd('write')
        ws = load_workbook(self.root / 'book.xlsx')['中国大陆']
        self.assertEqual(ws.cell(3, 11).value, S.LINK_EXCEPTION_PREFIX + '列表页按 positionId 定位，点岗位名展开 JD')
        self.run_cmd('link-exception', '--input', self.items([{**key0, "撤销": True}]), '--quote', '撤销')
        self.run_cmd('write')
        self.assertIsNone(load_workbook(self.root / 'book.xlsx')['中国大陆'].cell(3, 11).value)
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        row = next(r for r in master['sheets'][0]['rows'] if r['row'] == 4)
        S.set_cell(row, 11, S.LINK_EXCEPTION_PREFIX + '绕过脚本直接写')  # 绕过脚本直接改 JSON
        (self.root / '岗位交接/工作簿全量.json').write_text(json.dumps(master, ensure_ascii=False), encoding='utf-8')
        self.assertIn('没有用户确认记录', self.run_cmd('write', ok=False).stderr)

    def test_company_labels_live_in_own_group_file(self):
        """公司状态标签单独成组（一组一个文件）：岗位行只能用岗位组、公司行只能用公司组；跨组重名、组文件缺失都拒绝加载。"""
        p = self.run_cmd('propose', '--input', self.items([self.role(0, '已有岗位')]), ok=False)
        self.assertIn('不适用于岗位行', p.stderr)
        p = self.run_cmd('propose', '--input', self.items([{"sheet": "中国大陆", "company": "乙公司", "to": "岗位待筛选"}]), ok=False)
        self.assertIn('公司行只能用：待搜索、已有岗位', p.stderr)
        out = json.loads(self.run_cmd('propose', '--input', self.items([{"sheet": "中国大陆", "company": "乙公司", "to": "待搜索"}])).stdout)
        self.assertEqual(out['show_user_exactly'], ['中国大陆｜乙公司｜公司行：当前状态是「空」，更新状态是「待搜索」'])
        comp = self.root / '产品管理/company-labels.json'
        good = comp.read_text(encoding='utf-8')
        comp.write_text(json.dumps({"schemaVersion": 1, "rowKind": "company", "labels": COMPANY_LABELS + [{"value": "岗位待筛选"}]},
                                   ensure_ascii=False), encoding='utf-8')
        self.assertIn('多个标签组重复定义', self.run_cmd('check', ok=False).stderr)
        comp.unlink()
        self.assertIn('标签组文件不可读', self.run_cmd('check', ok=False).stderr)
        comp.write_text(good, encoding='utf-8')

    def test_every_company_keeps_exactly_one_company_row(self):
        """方案B：每家公司在每张地区表固定一行公司行；新增岗位不占用公司行；新公司自动补公司行（状态为空，须确认）。"""
        self.migrate()
        self.run_cmd('upsert', '--input', self.items([{"key": {"sheet": "中国大陆", "company": "乙公司", "link": "https://b.example/1", "title": "AI产品"},
                                                        "values": {"工作地点": "北京"}}]))
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        rows = [(r['row'], S.cell_value(r, 2), S.cell_value(r, 5), S.cell_value(r, 10)) for r in master['sheets'][0]['rows'][1:]]
        self.assertEqual(rows[-2:], [(7, '乙公司', None, '待搜索'), (8, '乙公司', 'AI产品', None)])
        out = json.loads(self.run_cmd('upsert', '--input', self.items([{"key": {"sheet": "中国大陆", "company": "丙公司", "link": "https://c.example/1", "title": "AI产品"},
                                                                        "values": {"工作地点": "上海"}}])).stdout)
        self.assertEqual(out['company_rows_created'], ['中国大陆｜丙公司'])
        master = json.loads((self.root / '岗位交接/工作簿全量.json').read_text(encoding='utf-8'))
        rows = [(r['row'], S.cell_value(r, 2), S.cell_value(r, 5)) for r in master['sheets'][0]['rows'][1:]]
        self.assertEqual(rows[-2:], [(9, '丙公司', None), (10, '丙公司', 'AI产品')])
        # 绕过脚本删掉公司行：门禁报缺公司行
        master['sheets'][0]['rows'] = [r for r in master['sheets'][0]['rows'] if not (S.cell_value(r, 2) == '甲公司' and S.cell_value(r, 5) is None)]
        S.refresh_digest(master)
        (self.root / '岗位交接/工作簿全量.json').write_text(json.dumps(master, ensure_ascii=False), encoding='utf-8')
        self.assertIn('中国大陆｜甲公司 缺少公司行', self.run_cmd('write', ok=False).stderr)

    def test_defer_binds_to_current_json(self):
        self.migrate()
        self.run_cmd('defer', '--quote', '下一轮再写')
        self.assertTrue(json.loads(self.run_cmd('pending').stdout)['deferred_by_user'])


if __name__ == '__main__':
    unittest.main()
